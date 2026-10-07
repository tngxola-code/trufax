from decimal import Decimal

from openpyxl import Workbook

from trufax.config import load_job, parse_job
from trufax.runner import run_job


def sheet_job(tmp_path, path, block: str = "", fields: str = "") -> object:
    fields = fields or ("    item: Item\n    amount: {src: Amount, type: decimal}")
    return parse_job(
        f"""
job: sheet-test
source:
  type: spreadsheet
  start_urls: ["{path}"]
  spreadsheet:
{block or "    header_row: 1"}
  fields:
{fields}
output: {{ formats: [csv] }}
""",
        tmp_path,
    )


def test_excel_example_with_totals_duplicates_and_row_numbers(examples, tmp_path):
    res = run_job(load_job(examples / "spreadsheet" / "price-list.yaml"), out_root=tmp_path)
    o = res.outcome
    assert o.counts == {"extracted": 5, "valid": 3, "invalid": 1, "duplicate": 1}
    total = [c for c in o.checks if "Total stock" in c.name]
    assert total and total[0].passed and "Stock row 9" in total[0].name
    drill = next(r for r in o.valid if r.data["sku"] == "DRL-1001")
    assert drill.provenance.locator == "Stock row 4"
    assert drill.data["price"] == Decimal(1249) and drill.data["in_stock"] is True
    dup = next(r for r in o.rejected if r.status == "DUPLICATE")
    assert "duplicate of Stock row 5" in dup.errors[0]


def test_csv_with_semicolons_and_a_wrong_total(tmp_path):
    f = tmp_path / "budget.csv"
    f.write_text("Item;Amount\nRates;1 200\nFees;300\nTotal;1 600\n", encoding="utf-8")
    block = (
        "    header_row: 1\n"
        '    totals: [{label_field: item, label_regex: "^Total", sum_fields: [amount]}]'
    )
    res = run_job(sheet_job(tmp_path, f, block), out_root=tmp_path / "o")
    assert [r.data["amount"] for r in res.outcome.valid] == [Decimal(1200), Decimal(300)]
    failed = [c for c in res.outcome.checks if not c.passed]
    assert failed and "rows sum to 1500" in failed[0].detail
    assert res.manifest.status == "REVIEW"
    assert res.outcome.valid[0].provenance.locator == "row 2"


def test_columns_by_index_and_sheet_by_name(tmp_path):
    wb = Workbook()
    wb.active.title = "Cover"
    ws = wb.create_sheet("Data")
    ws.append(["code", "label"])
    ws.append(["A1", "Alpha"])
    ws.append([None, None])  # blank rows are skipped
    ws.append(["B2", "Beta"])
    path = tmp_path / "book.xlsx"
    wb.save(path)
    fields = "    code: {src: 0}\n    label: {src: 1}"
    res = run_job(
        sheet_job(tmp_path, path, "    sheet: Data\n    header_row: 1", fields),
        out_root=tmp_path / "o",
    )
    assert [r.data["label"] for r in res.outcome.valid] == ["Alpha", "Beta"]
    assert res.outcome.valid[1].provenance.locator == "Data row 4"


def test_problems_become_warnings(tmp_path):
    f = tmp_path / "a.csv"
    f.write_text("Item,Amount\nx,1\n", encoding="utf-8")
    res = run_job(sheet_job(tmp_path, f, fields="    item: Missing"), out_root=tmp_path / "o")
    assert "column 'Missing' not found" in res.manifest.warnings[0]

    wb = Workbook()
    path = tmp_path / "b.xlsx"
    wb.save(path)
    res = run_job(sheet_job(tmp_path, path, "    sheet: Nope"), out_root=tmp_path / "o2")
    assert "sheet 'Nope' not found" in res.manifest.warnings[0]
