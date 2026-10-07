"""Generate examples/spreadsheet/supplier-price-list.xlsx (fictional data).

Run: python scripts/make_example_xlsx.py
"""

from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

OUT = Path(__file__).resolve().parents[1] / "examples" / "spreadsheet" / "supplier-price-list.xlsx"

ROWS = [
    ("DRL-1001", "Cordless Hand Drill", "Bosch", "Power tools", 1249.00, 14),
    ("MSR-0208", "Tape Measure 8m", "Stanley", "Measuring", 189.50, 102),
    ("PPE-0410", "Safety Gloves (pair)", None, "Safety", 64.99, 0),
    ("MSR-0208", "Tape Measure 8m", "Stanley", "Measuring", 189.50, 102),
    ("LGT-0050", "LED Work Light", "Ryobi", "Lighting", -449.00, 3),
]


def main() -> None:
    wb = Workbook()
    cover = wb.active
    cover.title = "Read me"
    cover["A1"] = "Example Supply Co. price list. Prices in ZAR, excluding VAT."
    ws = wb.create_sheet("Stock")
    ws.append(["Supplier price list", None, None, None, None, None])
    ws.append([f"Updated {date(2026, 10, 1).isoformat()}", None, None, None, None, None])
    ws.append(["SKU", "Product", "Brand", "Category", "Unit price", "Qty on hand"])
    for c in ws[3]:
        c.font = Font(bold=True)
    for row in ROWS:
        ws.append(list(row))
    ws.append([None, "Total stock", None, None, None, sum(r[5] for r in ROWS)])
    ws.append(["Notes: prices change monthly", None, None, None, None, None])
    wb.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
