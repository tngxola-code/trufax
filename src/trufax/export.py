"""Writing results: CSV, Excel, JSON and (optionally) a Google Sheet."""

from __future__ import annotations

import csv
import json
import os
from datetime import date
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .config import OutputConfig
from .records import PROVENANCE_COLUMNS, Record


def _plain(v: Any) -> Any:
    if isinstance(v, Decimal):
        return float(v) if v != v.to_integral_value() else int(v)
    if isinstance(v, date):
        return v.isoformat()
    return v


def rows_of(records: list[Record], include_provenance: bool, with_errors: bool = False):
    rows = []
    for r in records:
        row = {k: _plain(v) for k, v in r.flat(include_provenance).items()}
        if with_errors:
            row["_status"] = r.status
            row["_errors"] = " | ".join(r.errors)
        rows.append(row)
    return rows


def write_csv(path: Path, rows: list[dict], columns: list[str]) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=columns, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)


def write_xlsx(path: Path, sheets: dict[str, tuple[list[dict], list[str]]]) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    for title, (rows, columns) in sheets.items():
        ws = wb.create_sheet(title[:31])
        ws.append(columns)
        for c in ws[1]:
            c.font = Font(bold=True)
        for row in rows:
            ws.append([row.get(c) for c in columns])
        ws.freeze_panes = "A2"
        for i, col in enumerate(columns, 1):
            width = max([len(str(col))] + [len(str(r.get(col) or "")) for r in rows[:200]])
            ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 8), 60)
    wb.save(path)


def write_outputs(
    out_dir: Path, fields: list[str], valid: list[Record], rejected: list[Record], cfg: OutputConfig
) -> list[Path]:
    prov = [
        c
        for c in PROVENANCE_COLUMNS
        if c != "_ai_fields" or any(r.provenance.ai_fields for r in valid + rejected)
    ]
    columns = fields + (prov if cfg.include_provenance else [])
    rej_columns = fields + ["_status", "_errors"] + prov
    data = rows_of(valid, cfg.include_provenance)
    rej = rows_of(rejected, True, with_errors=True)
    written: list[Path] = []
    if "csv" in cfg.formats:
        write_csv(out_dir / "data.csv", data, columns)
        written.append(out_dir / "data.csv")
        if rej:
            write_csv(out_dir / "rejected.csv", rej, rej_columns)
            written.append(out_dir / "rejected.csv")
    if "xlsx" in cfg.formats:
        sheets = {"data": (data, columns)}
        if rej:
            sheets["rejected"] = (rej, rej_columns)
        write_xlsx(out_dir / "data.xlsx", sheets)
        written.append(out_dir / "data.xlsx")
    if "json" in cfg.formats:
        (out_dir / "data.json").write_text(json.dumps(data, indent=2, ensure_ascii=False))
        written.append(out_dir / "data.json")
    if cfg.google_sheet:
        push_to_sheet(cfg, data, columns)
    return written


def push_to_sheet(cfg: OutputConfig, rows: list[dict], columns: list[str]) -> None:
    """Replace a worksheet's contents. Needs: pip install 'trufax[sheets]' and a
    service-account JSON whose path is in the configured environment variable."""
    try:
        import gspread
    except ImportError as e:  # pragma: no cover - optional extra
        raise RuntimeError("Google Sheets output needs: pip install 'trufax[sheets]'") from e
    gs = cfg.google_sheet
    if gs is None:
        return
    creds = os.environ.get(gs.credentials_env)
    if not creds:
        raise RuntimeError(f"set {gs.credentials_env} to a service-account JSON file path")
    sh = gspread.service_account(filename=creds).open_by_key(gs.spreadsheet_id)
    try:
        ws = sh.worksheet(gs.worksheet)
        ws.clear()
    except gspread.WorksheetNotFound:
        ws = sh.add_worksheet(gs.worksheet, rows=len(rows) + 1, cols=len(columns))
    ws.update([columns] + [[r.get(c, "") for c in columns] for r in rows])
