"""CSV and Excel files: one record per row, columns matched by header or position.

Locators use real spreadsheet row numbers ("Stock row 14"), so a client can open the
file and find the exact cell. Printed total rows reconcile like PDF totals.
"""

from __future__ import annotations

import csv
import io
import re
from datetime import date, datetime, time
from pathlib import PurePosixPath
from urllib.parse import urlparse

from openpyxl import load_workbook

from ..config import SpreadsheetSettings
from ..fetch import Fetched, FetchError
from ..records import Record
from .base import Adapter
from .tabular import Row, norm_cell, records_with_totals, row_values

CSV_SUFFIXES = {".csv", ".tsv", ".txt"}
XLSX_SUFFIXES = {".xlsx", ".xlsm"}


def detect_format(doc: Fetched, cfg: SpreadsheetSettings) -> str:
    if cfg.format != "auto":
        return cfg.format
    suffix = PurePosixPath(urlparse(doc.final_url).path).suffix.lower()
    if suffix in XLSX_SUFFIXES or doc.content[:2] == b"PK":  # xlsx files are zip archives
        return "xlsx"
    if suffix in CSV_SUFFIXES or "csv" in doc.content_type:
        return "csv"
    raise FetchError(f"cannot tell whether {doc.final_url} is CSV or Excel; set format")


def cell_text(value: object) -> str:
    """Excel cells to text the transforms understand: dates as ISO, whole-number floats
    without '.0'."""
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time(0) else value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return norm_cell(value)


def read_rows(doc: Fetched, cfg: SpreadsheetSettings) -> tuple[str, list[list[str]]]:
    """Return (sheet label, all rows as text) for CSV or Excel content."""
    if detect_format(doc, cfg) == "csv":
        text = doc.content.decode(cfg.encoding, errors="replace")
        delimiter = cfg.delimiter
        if delimiter is None:
            try:
                delimiter = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|").delimiter
            except csv.Error:
                delimiter = ","
        rows = [
            [norm_cell(c) for c in r] for r in csv.reader(io.StringIO(text), delimiter=delimiter)
        ]
        return "row", rows
    book = load_workbook(io.BytesIO(doc.content), read_only=True, data_only=True)
    try:
        if isinstance(cfg.sheet, int):
            if cfg.sheet >= len(book.sheetnames):
                raise FetchError(f"sheet index {cfg.sheet} not found; sheets: {book.sheetnames}")
            ws = book[book.sheetnames[cfg.sheet]]
        else:
            if cfg.sheet not in book.sheetnames:
                raise FetchError(f"sheet '{cfg.sheet}' not found; sheets: {book.sheetnames}")
            ws = book[cfg.sheet]
        rows = [[cell_text(c) for c in r] for r in ws.iter_rows(values_only=True)]
        return f"{ws.title} row", rows
    finally:
        book.close()


class SpreadsheetAdapter(Adapter):
    def extract(self) -> list[Record]:
        cfg = self.job.source.spreadsheet
        assert cfg is not None  # set by SourceConfig for spreadsheet sources
        records: list[Record] = []
        for url in self.job.source.start_urls:
            if self.full(records):
                break
            try:
                doc = self.fetcher.get(url)
                label, rows = read_rows(doc, cfg)
            except FetchError as e:
                self.warnings.append(str(e))
                continue
            except (ValueError, OSError) as e:  # corrupt or unreadable workbook
                self.warnings.append(f"cannot read {url}: {e}")
                continue
            try:
                parsed = self._rows(rows, label, cfg)
            except ValueError as e:  # e.g. a configured column is not in the headers
                self.warnings.append(f"{doc.final_url}: {e}")
                continue
            if not parsed:
                self.warnings.append(f"no data rows in {doc.final_url}")
            records.extend(records_with_totals(self, parsed, doc, cfg.totals, len(records)))
        return records

    def _rows(self, rows: list[list[str]], label: str, cfg: SpreadsheetSettings) -> list[Row]:
        if len(rows) < cfg.header_row:
            return []
        headers = rows[cfg.header_row - 1]
        specs = self.job.fields
        skips = [re.compile(p) for p in cfg.skip_rows_matching]
        out: list[Row] = []
        for number, row in enumerate(rows[cfg.header_row :], start=cfg.header_row + 1):
            if not any(row):
                continue
            line = " ".join(c for c in row if c)
            if any(s.search(line) for s in skips):
                continue
            out.append((row_values(row, headers, specs), f"{label} {number}", line))
        return out
